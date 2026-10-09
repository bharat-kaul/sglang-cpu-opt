#!/usr/bin/env python3
"""Generate the DSv4 tree-pilot methodology playbook as a downloadable Word doc."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1F, 0x37, 0x64)
BLUE = RGBColor(0x2E, 0x5B, 0xA8)
GRAY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x1E, 0x7A, 0x3C)
OUT = "/data/nfs_home/bkaul/dsv4_pilot_playbook.docx"

d = Document()
base = d.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)


def h1(t):
    p = d.add_heading(t, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY
        r.font.size = Pt(15)
    return p


def h2(t):
    p = d.add_heading(t, level=2)
    for r in p.runs:
        r.font.color.rgb = BLUE
        r.font.size = Pt(12)
    return p


def para(t, color=None, bold=False, size=10.5, after=4):
    p = d.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(t, lvl=0, bold_lead=None):
    p = d.add_paragraph(style="List Bullet" if lvl == 0 else "List Bullet 2")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(t)
    else:
        p.add_run(t)
    p.paragraph_format.space_after = Pt(2)
    return p


def num(t, bold_lead=None):
    p = d.add_paragraph(style="List Number")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(t)
    else:
        p.add_run(t)
    p.paragraph_format.space_after = Pt(2)
    return p


# ---- title ----
t = d.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Model- & Platform-Agnostic Enablement + Optimization Playbook")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = NAVY
s = d.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("A reusable methodology for enabling AND optimizing any model on any accelerator "
              "(CPU shown; the same loop applies to other CPUs and to GPUs)")
r.font.size = Pt(12)
r.font.color.rgb = GRAY
dt = d.add_paragraph()
dt.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = dt.add_run("Platform and model constants are INPUTS, not assumptions. The concrete pilot instantiation "
               "(DeepSeek-V4-Flash on Intel Xeon / EMR) lives in the per-platform spec + roofline artifacts, not in this method.")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

# ---- overview ----
h1("Overview")
para("Goal: from a bare model config + checkpoint, enable AND optimize the model end-to-end on a target "
     "accelerator, driving final wall-time toward the sum of the per-op roofline floors — platform/runtime "
     "overhead squeezed out and every kept kernel at its ceiling. Nothing here is specific to one model or "
     "one chip: the platform constants, the op set, and the native-kernel library are INPUTS.", GRAY)
para("Two theses, one plugin (wire into the serving stack via its external/plugin mechanism — no fork):", bold=True)
bullet("ops already covered by the target's native high-throughput kernels (e.g. AMX/oneDNN on CPU, "
       "Tensor-Core/cuBLAS/Triton on GPU) are enabled by wiring + validation only (capability inheritance: "
       "dtype, prepack/layout, fused-MoE).", bold_lead="Thesis 1 (covered) — ")
bullet("genuinely novel ops get an AI-written, roofline-tuned native kernel (C/C++ for CPU, CUDA/Triton "
       "for GPU), human-gated.", bold_lead="Thesis 2 (novel) — ")
para("Scope discovery is the router (which ops are covered vs novel); the spine-leaf performance tree "
     "\u2014 the macro/meso/micro per-op roofline loop \u2014 is the engine. One pipeline, re-applied per "
     "(model, platform).", GRAY)
para("Deliverable contract (what the pipeline guarantees):", bold=True)
bullet("Every op / fused-op is AUTHORED as a native kernel \u2014 a C/C++ torch extension on CPU "
       "(CUDA/Triton on GPU). Nothing hot stays in eager python/torch.", bold_lead="1 \u2014 ")
bullet("Each authored kernel is TESTED standalone (parity + roofline) and ONLY THEN integrated into the "
       "serving stack's model behind the cosine-fingerprint gate \u2014 wired as a torch op, no fork "
       "(the pilot: SGLang external model package / PyTorch).", bold_lead="2 \u2014 ")
bullet("All end-to-end model runs execute through the serving-stack engine (the pilot: SGLang / "
       "sgl.Engine), not a bespoke loop \u2014 so wall-time reflects the real inference path.", bold_lead="3 \u2014 ")

# ---- governing two-phase ----
h1("Governing principle — two phases")
para("Per-op completeness FIRST, wall-time/ordering SECOND (a deliberate inversion of ROI-skip).", bold=True)
h2("Phase A — per-op peaking (ALL ops, including fused ops)")
num("Discover every op in the layer, including fused ops — a fused op is a first-class unit with its own roofline.")
num("Per op/fused-op: roofline + binding resource (BW / compute / latency) on the target platform (measured constants).")
num("Arithmetic-intensity improvability: can the binding resource be moved by a lever (increase M/batch, fusion)? Record the achievable ceiling.")
num("Implement the op in the target's native kernel language (C/C++ on CPU, CUDA/Triton on GPU) and optimize to a reasonable extent, driving it to its ceiling (BW peak / compute peak / latency floor).")
num("If the gap to roofline cannot be bridged by known techniques → surface to the user for review.")
num("Every op done this way — none skipped. Completion gate = every op/fused-op at its ceiling.")
h2("Phase B — wall-time + ordering (only after Phase A completes)")
num("Switch to wall-time; fix platform/runtime pathology first (CPU: OpenMP spin-wait/affinity, NUMA bind, allocator; GPU: launch/stream config, occupancy, memory pool; plus dispatch & batching) until kernel-domination clears the bar.")
num("Apply Amdahl ordering — prioritize/aggregate high-share work. Ordering is sequencing, not exclusion.")
para("Bounding valves that keep Phase A finite: \u201creasonable extent\u201d + \u201csurface if can\u2019t bridge.\u201d "
     "Latency-bound \u201cpeak\u201d = the dispatch/overhead floor (fuse/batch launches), not an AI target.", GRAY)

# ---- pipeline phases ----
h1("Pipeline phases (in order)")

h2("Phase 0 — Platform scan + intake + capacity/precision pre-flight (Gate 0)")
bullet("Run the uArch probe on the target device → measured machine constants: per-memory-domain BW, "
       "per-dtype compute peak, ridge = peak/BW, topology (CPU: NUMA/SNC; GPU: devices/interconnect), "
       "ISA/occupancy. No hardcoded constants — swap the spec file to retarget.")
bullet("Data-type audit from the real checkpoint (metadata-only first): {op, stored dtype, target "
       "compute dtype, kernel, inflation?, dtype-bridge, parity gate}. Never inflate a native low-bit "
       "checkpoint; sub-tile dtypes are moved low-bit and computed in the native compute dtype via fused dequant.")
bullet("Capacity pre-flight: resident weights + activations + KV/state vs per-domain memory floor → "
       "parallelism map (tp/EP; EP-first for MoE). On CPU tp>1 hurts latency-bound decode; on GPU pick "
       "tp/EP by memory fit + interconnect.")

h2("Phase 1 — Op scope discovery + decomposition (Gate 1)")
bullet("Enumerate every op (walk each layer index — heterogeneity is easy to miss); 100% compute "
       "coverage; anything unclassifiable is UNKNOWN → treated as a gap.")
bullet("Dependency-closure for non-compute substrate (KV pool, backends, novel non-attention families).")
bullet("Scope discovery assigns each op a lane: A = covered by a native kernel on the target (donor); "
       "B = missing on the target but a reference impl exists elsewhere (numerical oracle); "
       "C = missing, author from the math spec.")

h2("Phase 2 — Fusion pass → roofline (per op, target platform)")
bullet("Fusion analysis first (vertical / horizontal / epilogue). For each candidate: consult the "
       "GPU/XPU implementation where available — runnable fused code is a numerical ORACLE; a "
       "performance claim without code is a HINT only — OR derive the benefit MATHEMATICALLY on the "
       "roofline (bytes saved + AI before→after + L2-resident gate) where no reference exists. Absence "
       "of a GPU reference is not a reason to skip (CPU is more BW-bound, so a CPU-only fusion may pay).")
bullet("Map each fusion: COVERED (donor fused kernel) / NEW-C++ (→ kernel authoring) / SKIP (record why).")
bullet("Per-op and per-fused-op roofline: FLOPs, weight+activation+KV bytes, arithmetic intensity, "
       "binding resource, achievable ceiling, and the native-kernel implementation lane. Phase-split prefill vs decode.")

h2("Phase A / Phase B — execute per the governing principle above")

# ---- progressive wall-time tracking ----
h1("Progressive wall-time tracking (publication ledger)")
para("Each optimization is proven end-to-end, not just in a microbench. We keep a running "
     "wall-time ledger from the unoptimized baseline through every integrated op, so the final "
     "publication shows a monotonic, correctness-gated speed-up curve.", bold=True)
para("Use the CHEAPEST representative proxy at every altitude, wherever feasible, and recover the "
     "full-scale number ANALYTICALLY \u2014 validated by ONE full run at the end: (a) op optimization uses the "
     "isolated-op microbench at real op dims (and reduced context S where the op scales trivially in S); "
     "(b) wall-time uses the depth-reduced model with linear depth-scaling to full N. Proxies move us "
     "faster; the single full-depth run at phase end confirms the analytical scaling held.", GRAY)
para("Ledger artifact: plugin/validate/results/wall_time_progress.json (baseline row + one row per "
     "op, filled after each integration).", GRAY)
h2("Parallel authoring, serial integration, per-pass record")
para("Ops are independent work items (distinct kernel, distinct microbench), so AUTHOR THEM IN PARALLEL "
     "(one op per idle node) and QUEUE the finished best-of versions. INTEGRATION stays serial: dequeue "
     "one op at a time, cosine-gate, re-run the wall-time ledger. Only the integrated re-profile is serial.", GRAY)
bullet("Per op, record EVERY optimization pass (before/after): approach, off-machine-peak and vs-reference "
       "per M, correctness (cosine / set-match), kept?, commit. The final KEPT pass is the best-of, reported "
       "against the roofline (at-ceiling, or surfaced-as-plateaued with the residual-gap cause).")
bullet("Per-op pass ledgers: plugin/validate/results/op_passes/<op>.json (schema in _schema.json).")
bullet("Torch/donor fallback registry: every point where the kept best-of falls back to torch (or a "
       "donor) instead of the custom path is recorded with its reason + limit justification "
       "(plugin/validate/results/torch_fallbacks.json). A fallback is legitimate only if strictly best-of "
       "(no regression at that point) OR at the measured limit with no lever left. The WHOLE registry is "
       "SURFACED FOR REVIEW at the end of the optimization + integration phase.")
h2("The loop (repeat per op, in Phase-A order)")
num("Baseline wall time FIRST: run the model end-to-end (depth proxy, deterministic-dummy weights) and "
    "record the steady-state median latency at each M (step 0 = the reference; never time the cold step).",
    bold_lead="0 \u2014 ")
num("Optimize the op standalone to its machine-peak ceiling (author \u2192 measure \u2192 iterate \u2192 keep "
    "best-of), then surface-and-stop when levers are exhausted.", bold_lead="1 \u2014 ")
num("Integrate the kept kernel into the model behind a COSINE FINGERPRINT gate: per-layer + per-op "
    "cosine vs a trusted reference (eager / GPU / HF) on the SAME dummy weights must stay \u2265 0.9999 "
    "(set-match for top-k ops). The first layer/op below threshold localizes a break \u2014 do not proceed until it passes.",
    bold_lead="2 \u2014 ")
num("Re-run the model end-to-end (dummy weights, same M sweep) and record the new wall time + the delta "
    "vs the previous row and the cumulative speed-up vs baseline.", bold_lead="3 \u2014 ")
num("Repeat for the next op. The ledger accumulates to the final published curve: baseline \u2192 "
    "op-by-op \u2192 fully-optimized.", bold_lead="4 \u2014 ")
h2("Depth-proxy timing + analytical scaling to full depth")
para("Perf runs use DUMMY weights (no real load) AND a depth-reduced proxy. When the layer stack is a "
     "repeated archetype (homogeneous transformer \u2014 the common case), end-to-end wall time is LINEAR in "
     "depth, t(L) = t_fixed + L\u00b7t_layer, so a few-layer proxy is representative (pilot: 4 layers vs the "
     "full N). For HETEROGENEOUS stacks, include each distinct archetype at least once and scale per-archetype.", GRAY)
bullet("Calibrate ONCE at baseline from two proxy depths p and q: t_layer = (t_q \u2212 t_p)/(q \u2212 p) and "
       "t_fixed = t_p \u2212 p\u00b7t_layer, where t_fixed is the non-layer cost (embed + output head + "
       "framework/sampling). (pilot p=4, q=8.)")
bullet("Each iteration records BOTH: (a) the MEASURED proxy wall time, and (b) the full-depth wall time "
       "derived ANALYTICALLY \u2014 t_layer(iter) = (t_proxy(iter) \u2212 t_fixed)/p, predicted full = "
       "t_fixed + N\u00b7t_layer(iter). Report proxy and full-scaled speed-up curves side by side "
       "(they differ because the fixed overhead does not scale).")
bullet("t_fixed is config-invariant across per-layer optimizations; recalibrate only if a NON-layer op "
       "(e.g. the output head) is optimized.")
bullet("END VALIDATION: after all ops are integrated, do ONE full-depth dummy run (all N layers) and "
       "compare the MEASURED full-depth wall time against the analytically predicted value \u2014 small error "
       "confirms the whole depth-scaled progression.")
h2("Framework / dispatch / OpenMP overhead \u2014 the proxy studies these too (Phase B on the proxy)")
para("The depth proxy is the right instrument for system overhead, not just op math: the depth-independent "
     "t_fixed term IS the framework/dispatch/sampling cost, and per-op dispatch + OpenMP-barrier overhead "
     "lives in t_layer and is incurred every layer \u2014 both are fully present at depth 4. Global OMP "
     "pathologies are depth-invariant, so the proxy reproduces them at a fraction of the cost.", GRAY)
num("Finish ALL ops/fused-ops to their roofline ceiling and integrate them into the proxy (end of Phase A); "
    "record the fully-integrated proxy + full-scaled wall time.", bold_lead="Order \u2014 ")
num("THEN attack framework/dispatch/OMP overhead on that integrated proxy (this is Phase B). Bit-exact "
    "gate: an overhead fix must not change tokens.", bold_lead="")
num("Show the wall-time improvement from the overhead work (proxy + full-scaled), SURFACE it, and only "
    "THEN launch the single full-depth run.", bold_lead="")
bullet("Bind threads ONCE at init \u2014 never per-forward torch.set_num_threads (it rebuilds the OpenMP pool "
       "and LOSES CPU affinity \u2192 unpinned contention inflates the whole forward uniformly; invisible per-op).")
bullet("Collapse many tiny torch ops into fewer/bigger parallel regions \u2014 at N threads each tiny op pays "
       "a barrier, so ~100 ops/layer is a ~10-20x dispatch tax; fuse into one parallel region.")
bullet("Clear the usual systemic levers: OMP spin-wait (OMP_WAIT_POLICY), allocator (tcmalloc/jemalloc), "
       "NUMA first-touch into the rank's own domain, thread-cap, CPU-freq/cold-start; remove dtype-cast and "
       "dispatch churn on the hot path. A global-config pathology masquerades as uniform per-op cost \u2014 "
       "clear it before trusting any per-op ranking.")
h2("MoE mode-collapse \u2014 the input tensor matters")
para("The dummy-weight run must exercise a REPRESENTATIVE set of routed experts. With tiny random gate "
     "weights, a non-diverse input (e.g. one hidden vector broadcast across the batch) collapses routing "
     "to the same top-k experts every token \u2014 only ~topk of the E experts run, so the measured MoE "
     "wall time is unrepresentatively low and the roofline's expert traffic is never incurred.", GRAY)
bullet("Build the input with plugin/validate/moe_balanced_input.build_verified_input(M, gate_weight): "
       "per-token i.i.d. hidden states (NOT a broadcast), falling back to round-robin gate-row alignment.")
bullet("VERIFY (not assume) dispersion: realized distinct-experts must track the uniform expectation "
       "E\u00b7(1\u2212(1\u22121/E)^(topk\u00b7M)) and no expert may absorb an outsized share \u2014 a collapse gate "
       "asserts this before the timing is trusted. (pilot E=256: checked M=8\u219243, 16\u219281, 32\u2192138, "
       "64\u2192209 distinct vs expected 43.8/80.2/135.3/199; the broadcast trap collapses to 6 and is rejected.)")
bullet("Instrument the real router during the E2E run to confirm the same spread in-engine (per-forward "
       "distinct-expert count), so the proxy input is validated against the actual model path.")

h2("Phase 6 — Validate, certify, deliver")
bullet("Validate-by-run: confirm the predicted class/ceiling in-engine (dummy weights for perf, full "
       "depth, target backend). Mismatch → the paper model omitted a cost; update it.")
bullet("Accuracy finale: real-weight per-token parity vs a trusted reference → task accuracy "
       "(e.g. gsm8k) → reference-vs-target cross-check sign-off.")
bullet("Provenance record (op→kernel→donor + parity gates) and a peer-relative check where an already-"
       "shipped model shares the same kernels (efficiency within tolerance); ship as a plugin override (no fork).")

# ---- correctness toolkit ----
h1("Correctness debugging toolkit (localize cheaply before real-weight runs)")
bullet("Deterministic-dummy weights: (a) skip the real checkpoint load for fast perf iteration; "
       "(b) make the target dummy bit-identical to the reference dummy so a per-layer diff is valid "
       "without real weights.")
bullet("Layer-by-layer cosine fingerprint: capture a per-layer reference, run the target with the same dummy "
       "weights, diff per-layer hidden states by cosine similarity; the first layer below ~0.999 localizes the break.")
bullet("Per-op fingerprint (shape + norm + cosine/checksum) to pinpoint the op within the bad layer.")
bullet("Isolation switches to bisect novel ops (e.g. bypass sparse selection → dense, force-dense MoE, swap "
       "kernel ↔ reference).")
bullet("Reduces-to-identity checks (e.g. sparse attend → exact dense at top-k=all) and low-bit dtype-bridge "
       "parity vs an independent dequant oracle.")
bullet("Escalation order: L0 coherence → per-layer cosine (dummy, no load) → per-op fingerprint + "
       "isolation → component/reduces-to-identity parity → real-weight per-token parity → task accuracy "
       "→ reference-vs-target sign-off.")

# ---- shape-capture audit ----
h1("Shape & semantics provenance — ground kernels in the PUBLISHED reference (human-auditable)")
para("A green parity microbench is SELF-CONSISTENT only: it validates op MATH at shapes the AUTHOR chose. "
     "Shape-generic ops (softmax-pool, flash-attention) pass cos=1.0 at ANY dims, so a hallucinated shape "
     "OR a missing semantic (an overlap window, an attention sink, MQA-vs-MHA) still passes. Ground every "
     "kernel's shapes AND semantics in the model author's PUBLISHED reference, not in author-chosen shapes "
     "and not in inference from the config alone.", bold=True)
h2("The method (model-agnostic, in order)")
num("PIN the published reference model definition (the author's own model.py / modeling_*.py) at a fixed "
    "revision. This is the architecture a human reviewer already trusts \u2014 the audit anchor.", bold_lead="1 \u2014 ")
num("Derive, per op, a CONTRACT from it: shapes as f(config, per-layer variant, phase, TP, batch) AND the "
    "SEMANTICS (fusions, gating, sinks, overlap windows, MQA/MHA, two-of-a-kind modules). Cite the source line.",
    bold_lead="2 \u2014 ")
num("BIND numeric constants from the real checkpoint config.json \u2014 the reference's defaults are often a "
    "TOY example (e.g. fewer layers/experts). CROSS-CHECK config vs reference defaults and reconcile any "
    "divergence; this also audits a regenerated/proxy config.", bold_lead="3 \u2014 ")
num("SWEEP the contract over the REAL variation axes \u2014 per-layer variants (e.g. per-layer compression "
    "ratios), prefill vs decode, TP, batch M \u2014 and assert each shape is VALID and EQUALS the kernel's "
    "contract (equality, not mere internal validity).", bold_lead="4 \u2014 ")
num("Port the published forward as the NUMERICAL ORACLE (exact parity, not a re-derived reference) so a "
    "SEMANTIC error, not just a shape error, is caught.", bold_lead="5 \u2014 ")
num("RUNTIME-CAPTURE (module-level now, full-forward at integration) to CONFIRM the serving stack's actual "
    "layout matches the contract \u2014 the serving path may repack/absorb/shard differently from the reference.",
    bold_lead="6 \u2014 ")
para("Why it survives human audit: every shape and semantic traces to a line in the PUBLISHED model + the "
     "real config; the sweep, validity and equality checks are mechanical; the oracle gives exact numbers. "
     "The reviewer checks FACTS and CITATIONS, not the agent's reading.", GRAY)
para("Anti-patterns (each is a silent-failure trap): trusting a parity microbench at author-chosen shapes; "
     "inferring dims from config alone (misses semantics \u2014 overlap, sink, MQA); treating the reference's "
     "toy defaults as the real constants; and the cardinal sin \u2014 RE-GUESSING a shape after finding one "
     "wrong instead of capturing it from the published model.", GRAY)

# ---- kernel-algorithm provenance ----
h1("Kernel-algorithm provenance — trace EVERY authored kernel to the published reference (REQUIRED before impl review)")
para("A kernel is trusted only when its ALGORITHM traces to an EXTERNAL reference \u2014 the published model\u2019s "
     "op definition or primitive call site, plus the authoritative numerical oracle \u2014 with an explicit "
     "FAITHFUL / FRAGMENT / DIVERGENT disposition and ENUMERATED gaps. A green parity microbench proves the "
     "MATH at the author\u2019s shapes; it does NOT prove the kernel implements the MODEL\u2019s op: a stage-faithful "
     "FRAGMENT (pool core without the projections/RoPE/Hadamard/quant/decode-state around it) passes cos=1.0 "
     "while silently omitting most of the subgraph. Establish this BEFORE reviewing or re-authoring "
     "implementations. Model-agnostic.", bold=True)
h2("The method (in order)")
num("PIN the published reference and LOCATE the op: its module/function definition OR its call site + "
    "signature. For a PRIMITIVE imported from a kernel module (CUDA/TileLang, not in model source), take the "
    "integration contract from the call site and the NUMERICS from the published torch fallback/oracle \u2014 and "
    "say so (provenance-depth caveat).", bold_lead="1 \u2014 ")
num("MAP the kernel to a STAGE of that reference and CLASSIFY: FAITHFUL (implements the stage), FRAGMENT "
    "(faithful core \u2014 then ENUMERATE every missing sub-op: projections, norm, RoPE, Hadamard, quant, overlap "
    "window, decode-state buffers, masking, gather), or DIVERGENT (semantics differ \u2014 stop).",
    bold_lead="2 \u2014 ")
num("CITE the numerical ORACLE and the microbench EVIDENCE. Correctness is MICROBENCH scope (current-target/"
    "E2E verification PENDING); record the kept-pass REVISION (UNRESOLVED if the record commit is pending). "
    "Mark a SURFACED / non-production kernel as such (e.g. a scalar kernel that loses to a donor path).",
    bold_lead="3 \u2014 ")
num("RECORD the determination in a structured artifact (one entry per kernel: external_reference, oracle, "
    "stage, disposition, unmodeled_gaps, validation). CROSS-LIST every gap with the cost model\u2019s "
    "EXPLICITLY-UNMODELED items so the two artifacts agree.", bold_lead="4 \u2014 ")
num("GATE: no implementation review, re-authoring, or performance trust for a kernel until it has FAITHFUL "
    "or FRAGMENT(+enumerated gaps) EXTERNAL provenance recorded. Provenance is a prerequisite, not a "
    "by-product of optimization.", bold_lead="5 \u2014 ")
para("Why it survives human audit: the reviewer opens the referenced module line, the oracle, and the "
     "record; the stage/disposition/gaps are explicit and external. Anti-patterns: a parity microbench taken "
     "as proof the kernel implements the model\u2019s op; a stage-faithful FRAGMENT reported as the full subgraph; "
     "gaps left unspoken; citing a locally re-derived reference as the authority instead of the published "
     "module; and re-authoring before provenance exists.", GRAY)
h2("Kernel \u2194 cost-model reconciliation (the two artifacts must not be disconnected)")
para("The kernel provenance (what each kernel COMPUTES) and the roofline cost model (what each op is COSTED "
     "as) are two artifacts describing the SAME ops. They must RECONCILE or the roofline is costing something "
     "the kernels don\u2019t compute (or vice-versa). Make this a fail-closed cross-check, not a hope.", bold=True)
num("MAP every authored kernel to the cost row(s) that cost it; assert each named cost row EXISTS in the "
    "model\u2019s op inventory (a rename/drop on either side must BREAK the check).", bold_lead="1 \u2014 ")
num("DISPOSITION every FRAGMENT gap against the cost model: each gap is modeled by ANOTHER cost row, OR is "
    "an EXPLICITLY-UNMODELED tracker item, OR is a caller op \u2014 never nothing. A gap that is none of these is "
    "a DISCONNECT.", bold_lead="2 \u2014 ")
num("CHECK operand/shape/dtype agreement: the cost row\u2019s operands and FLOP structure must match the kernel\u2019s "
    "I/O contract (the same benchmark contract the provenance cites as oracle).", bold_lead="3 \u2014 ")
num("NAME dual paths explicitly: if the costed path differs from an authored kernel (e.g. the model costs a "
    "DONOR flash while an authored scalar kernel is SURFACED), record both so they are not conflated.",
    bold_lead="4 \u2014 ")
num("ENCODE it as a self-test the report emission depends on: load the provenance record + the op inventory "
    "+ the unmodeled tracker and assert 1\u20134; fail closed on a missing record.", bold_lead="5 \u2014 ")

# ---- roofline accounting discipline ----
h1("Roofline accounting discipline — REFERENCE-CONFORMANT, fail-closed (self-consistency is NOT the gate)")
para("An analytical roofline is trustworthy only when every FLOP/byte/dtype/count traces to an EXTERNAL "
     "reference (above) AND the accounting obeys the rules below. Encode the rules as CONFORMANCE TESTS that "
     "GATE report generation and FAIL CLOSED \u2014 a model that fails a check, or whose evidence is missing, is "
     "not published. Internal consistency (a green self-test) is necessary but NEVER sufficient: it proves "
     "the artifact agrees with itself, not with the model. These rules are model-agnostic.", bold=True)
h2("Knob\u2013cost coupling discovery — the generative meta-discipline (KV\u00d7M is ONE instance, not the lesson)")
para("The deepest lesson is not any single rule below \u2014 it is the METHOD that generates them: a cost model "
     "is cost = f(knobs), and almost every accounting error is a DROPPED COUPLING between a knob and a cost "
     "term (a term wrongly treated as independent of an axis it actually moves with). KV\u00d7M is one such "
     "coupling; the goal is a procedure that DISCOVERS the others by construction so a reviewer never has to.",
     bold=True)
para("Procedure (do this before ranking anything): (1) ENUMERATE the independent axes \u2014 workload knobs "
     "(batch M, context/seq S, per-layer sparsity/compression ratio, window, top-k k & expert count E, "
     "speculative depth, prefill-vs-decode) and hardware knobs (BW, compute peak PER dtype, cache capacity, "
     "NUMA/TP domain count). (2) Build a KNOB\u00d7COST-TERM matrix: for every cost term (each operand\u2019s bytes, "
     "each op\u2019s FLOPs, each invocation count) DECLARE its dependence on EACH axis and its functional form \u2014 "
     "flat / linear / saturating / piecewise / a PRODUCT of two knobs. (3) Hunt the knob-PAIR products "
     "explicitly \u2014 those are where naive models silently drop a cross-term. (4) SWEEP each axis and assert "
     "the declared form holds (linear stays linear, flat stays flat); a cross-term that should appear and "
     "doesn\u2019t, or appears where it shouldn\u2019t, is the bug.", GRAY)
para("Known couplings (INSTANCES of the matrix \u2014 extend per model/op/HW; each detailed as a rule below):",
     GRAY)
bullet("state \u00d7 batch: per-request state (KV/activations) bytes are LINEAR in M; weights are FLAT in M.",
       bold_lead="KV\u00d7M \u2014 ")
bullet("distinct-working-set \u00d7 batch: distinct routed experts streamed = E\u00b7(1\u2212(1\u2212k/E)^M) \u2014 SATURATING in M, "
       "not linear; dense/attention weights amortize flat over the batch.")
bullet("read-set = context \u00d7 sparsity-ratio: the positions an op reads are seq/ratio (+ window), summed over "
       "the per-layer ratios \u2014 a PRODUCT of two knobs, not the full sequence.")
bullet("invocation \u00d7 sequence: boundary-triggered ops fire seq/cadence times \u2014 amortized calls/step, a "
       "ratio of two knobs, not one-per-token.")
bullet("dtype \u00d7 compute-resource: the compute peak is SELECTED by the compute dtype (matrix-engine vs "
       "vector); storage dtype is a separate axis. A precision change that does not move the time is a "
       "dropped coupling.")
bullet("ceiling = per-domain-BW \u00d7 domains (with NUMA-local sharding): aggregate bandwidth scales with usable "
       "domains ONLY when weights are sharded one-rank-per-domain; an unsharded replica is capped at ONE "
       "domain. Capacity-fit \u00d7 head/expert-divisibility bounds the usable domains below the physical count.")
para("When a review finds a new miss, add the coupling to the matrix and a sweep-test for its form \u2014 that is "
     "how the playbook learns a connection once and never re-discovers it by hand.", GRAY)
bullet("ONE tensor inventory: capacity AND per-op cost derive from the SAME op list \u2014 never a second "
       "hardcoded capacity formula that can drift from the op dimensions.")
bullet("DECLARE the workload: independent requests vs shared-prefix, what is reused, dtype per tensor. An "
       "optimistic model is acceptable only if the assumption is stated.")
bullet("STATE \u00d7 BATCH is the error that keeps biting: per-request state (KV, activations, scratch) scales "
       "with batch M; weights stream once per step (M-independent). TEST: independent-request bytes are "
       "LINEAR in M, weight bytes are FLAT in M. A state term that is accidentally M-independent (or a weight "
       "term that scales with M) is an accounting bug.", bold_lead="KV\u00d7M \u2014 ")
bullet("DISTINCT working sets can grow NON-linearly with M: a routed/MoE layer streams the DISTINCT experts "
       "touched by the batch = E\u00b7(1\u2212(1\u2212k/E)^M) (distinct top-k of E over M tokens), not k\u00b7M and not E. Use "
       "the with-REPLACEMENT-free form; TEST it equals k at M=1 and saturates toward E.")
bullet("The per-op READ SET is the reference's ACTUAL access pattern, not the full sequence: sliding-window "
       "/ sparse / compressed attention reads window+selected positions summed over the per-layer variants "
       "\u2014 never full-context on every layer, and never a full pass plus a separate selection pass if the "
       "reference fuses them. Derive positions from the config's per-layer ratios.")
bullet("ENUMERATE EVERY STREAM an op touches from its I/O contract: a fused op reads ALL its operands (a "
       "softmax-pool reads BOTH values AND scores), writes its outputs, and reads shared constants once. "
       "Dropping a stream silently understates traffic. Tag each operand input/output/weight/state + dtype.")
bullet("SHARED vs PER-REQUEST operands scale differently with M: a tensor read once per call is M-flat; a "
       "per-request tensor is \u00d7M. Mis-tagging one is an M-scaling bug (it hides until a batch sweep).")
bullet("BOUNDARY-TRIGGERED ops are AMORTIZED by cadence, not counted \u00d7full-length: an op firing every r "
       "tokens contributes r-fractional calls/step; an overlap window changes the pooled span. Model the "
       "amortized calls/step, not one call per token.")
bullet("FLOPs include REDUCTION ADDS, not just multiplies: a k-term weighted sum is k multiplies + (k\u22121) "
       "adds; a matmul is 2\u00b7M\u00b7N\u00b7K. Counting only multiplies understates compute.")
bullet("STORAGE dtype \u2260 COMPUTE dtype. Storage bytes (incl scale/padding) come from the checkpoint header; "
       "the COMPUTE RESOURCE/peak is selected BY the compute dtype (matrix-engine for bf16/int8, vector unit "
       "for fp32, \u2026) from a cited datasheet. TEST: a precision change MUST change the computed time; a dtype "
       "label that leaves the time unchanged is not modeled.", bold_lead="dtype \u2014 ")
bullet("FUSION removes round-trips, not weights: a fused op counts ALL its matrices (a SwiGLU expert is "
       "three), and a flash/fused op carries NO intermediate-score DRAM. TEST both structurally.")
bullet("RIDGE crossing solves the REAL arithmetic intensity (weights + activations) and returns "
       "\u201cno crossing\u201d when AI saturates below the ridge \u2014 not a weight-only approximation.")
bullet("Per-layer INVOCATION counts come from the config (per-layer variants/phases), not a uniform "
       "\u00d7num_layers. Verify against the model; print per-call cost, calls/step, and per-step cost.")
bullet("Include quantization metadata (scale/padding bytes) in traffic AND capacity; low-bit storage is "
       "not native low-bit compute.")
bullet("SWEEP M=1..64 in BOTH the roofline and the measurement \u2014 never a single point. The analytical "
       "ideal and the measured median must be reported at EVERY M (1/8/16/32/64); a single-M number cannot "
       "show the plateau or the regime. Expect the measured off-ceiling to CONVERGE toward the achievable "
       "BW wall as M grows while the small-M points are dispatch/overhead-bound \u2014 that convergence IS the "
       "best-case-vs-roofline evidence. Measure replicated (median of >=3), threads bound once, one NUMA "
       "domain, on the roofline's measured-reference node class.", bold_lead="M-sweep \u2014 ")
bullet("MEASUREMENTS are VALIDATED against their raw record, not asserted: a cited speedup/correctness must "
       "be READ from a structured result record (load it; fail closed if absent). KERNEL revision and "
       "RESULT-RECORD revision are SEPARATE fields; the published quantity must match the specific KEPT "
       "result entry at the stated coordinate. Report the recorded correctness field VERBATIM under a fixed "
       "attribution (\u2018historical author-reported microbench result; test scope UNVERIFIED; E2E verification "
       "PENDING\u2019) \u2014 do NOT infer a positive match, metric, or certification from free text (that would "
       "\u2018certify\u2019 an absent or FAILing record): absent => \u2018unrecorded\u2019, a reported failure is retained as a "
       "failure. WITHHOLD a speedup whose record ratio is at a superseded coordinate or whose revision is "
       "unresolved. Never hand-type a certification string the generator does not read back; regression-test "
       "that missing / failing / failed-cosine evidence generates NO certification.",
       bold_lead="measurement \u2014 ")
bullet("Latency rows are MEASURED (validated against a record), justified-analytical, or "
       "EXPLICITLY-UNMODELED \u2014 never invented floors; an absolute no record contains is WITHHELD, not "
       "populated; a single observed point is not a batch sweep (render only at its coordinate).")
bullet("Record hardware-profile PROVENANCE (node, memory type, clock); never mix measurements across "
       "configurations. The ideal roofline is an optimization TARGET, not an achievability claim.")
bullet("Report latency and useful throughput as PRIMARY; track useful-vs-executed FLOPs/traffic "
       "separately; use distance-from-roof only DIAGNOSTICALLY (a fused kernel can be faster at LOWER "
       "achieved bandwidth). Emit a roofline-VS-observation join per authored op; draw NO causal "
       "(overhead-vs-compute) or donor-routing conclusion from distance alone.")
bullet("PROCESS: read external reviews/commits that touch your area BEFORE building on them \u2014 a rebase "
       "is not a read.")
h2("Open-item disposition (never silently approximate)")
para("Every quantity in the model is exactly one of: MODELED (derived, cite the source line), MEASURED "
     "(microbench), or EXPLICITLY-UNMODELED (declared + tracked + surfaced at review with a reason and a "
     "plan). No silent placeholders or invented constants. Unmodeled items are listed in the review, never "
     "hidden inside a number. Prefer MODELED when the reference gives a formula; fall back to "
     "EXPLICITLY-UNMODELED only when it genuinely cannot be derived yet.", GRAY)

# ---- external-reference provenance (the auditability discipline) ----
h1("External-reference provenance — everything is EXTERNALLY referenced (never self-referenced)")
para("The single discipline behind every correctness gate: EVERY op semantic, shape, number, dtype, "
     "invocation count, kernel claim and measurement must trace to an EXTERNAL AUTHORITY that a reviewer "
     "can open and check \u2014 never to the artifact's own assertion, and never to a plausible reconstruction "
     "from memory. External-referenced = auditable. Self-referenced = unfalsifiable. A model that only "
     "agrees with itself is self-consistent, not correct. This is model-agnostic: it holds for any model, "
     "op, kernel, or hardware.", bold=True)
para("WHY THIS KEEPS BITING (the recurring failure modes \u2014 each is a self-reference leaking in): a number "
     "modeled from plausibility instead of transcribed from the source; a LABEL changed (dtype/annotation) "
     "while the computed consequence is unchanged, and a test that checks the label; a gate that validates "
     "the artifact against itself and FAILS OPEN when evidence is absent; patching only the counterexample a "
     "reviewer cited instead of the whole error CLASS; a derived byte/FLOP invented instead of enumerated "
     "from the op's input/output contract; and an absolute measurement rendered as certified when no raw "
     "record contains it. All five were caught by an external reviewer that the self-tests could not catch.",
     GRAY)

h2("The authoritative external references (open one of these, or it is UNVERIFIED)")
bullet("SEMANTICS + SHAPES: the author's PUBLISHED model definition (model.py / modeling_*.py) PINNED at a "
       "fixed revision. Cite the source line for each fusion, gate, sink, overlap window, MQA/MHA, per-layer "
       "variant. The published forward is also the numerical ORACLE.")
bullet("STORAGE DTYPE: the real CHECKPOINT HEADER (e.g. safetensors header via range request) \u2014 never "
       "inferred. Storage dtype is a SEPARATE fact from compute dtype (low-bit storage \u2260 low-bit compute).")
bullet("NUMERIC CONSTANTS: the real config.json (counts, ratios, dims) \u2014 cross-checked against the "
       "reference's defaults, which are often a TOY example.")
bullet("DERIVED BYTES/FLOPs: the op's actual INPUT/OUTPUT CONTRACT \u2014 the reference module signature OR the "
       "benchmark call signature. Enumerate operands from it; tag each (input/output/weight/state, dtype, "
       "per-request vs shared). Do not invent or drop operands.")
bullet("COMPUTE RESOURCE: selected BY the compute dtype (bf16\u2192matrix-engine peak, fp32\u2192vector peak, \u2026) "
       "from a cited HW datasheet/probe. A dtype label that does not change the computed time is not modeled.")
bullet("MEASUREMENTS: an AUDITABLE raw RESULT RECORD (result file + kernel/bench revision/commit), not a "
       "filename or a model SHA. Cite what the record actually certifies (coordinate, ownership, dtype, "
       "correctness, speedup). If no record holds the quantity (e.g. an absolute latency), WITHHOLD it as "
       "unverified \u2014 do not invent provenance and do not rerun expensive jobs just to populate a table.")

h2("The rules that make it auditable (model-agnostic)")
num("PROVENANCE TRIPLE on every published quantity: (value, cited external source, derivation). No "
    "citation \u21d2 EXPLICITLY-UNMODELED, not MODELED. Generalize this from kernel shapes to EVERY number: "
    "bytes, FLOPs, dtypes, counts, invocation cadence, capacity.", bold_lead="1 \u2014 ")
num("ASSERT THE CONSEQUENCE, never the label. A conformance test must check the downstream COMPUTED output "
    "(rendered time, selected peak, byte total, emitted row), not a field/annotation. Counter-test: "
    "re-injecting the OLD behavior must FAIL the test. A test that can pass while the output is unchanged is "
    "not a gate.", bold_lead="2 \u2014 ")
num("REFERENCE-CONFORMANCE, FAIL-CLOSED. The gate compares the artifact to the external reference, not to "
    "itself, and ABORTS emission on missing / malformed / invalid / incomplete evidence. Exercise the REAL "
    "validator with invalid inputs (wrong enum, missing identity, dropped coverage, duplicate) \u2014 never a "
    "re-implementation of the check.", bold_lead="3 \u2014 ")
num("FIX THE CLASS, not the counterexample. When a review finds one defect, enumerate EVERY sibling of the "
    "same type and add a class-covering test before resubmitting. Reviewers sample classes; clearing one "
    "instance invites the same finding on the next instance.", bold_lead="4 \u2014 ")
num("OBSERVATION \u2260 IDEAL TARGET, and a single point \u2260 a sweep. Render a measurement ONLY at its sourced "
    "coordinate; never broadcast one constant across a batch sweep; keep measured observations visually "
    "distinct from analytical targets; withhold any absolute that no record contains.", bold_lead="5 \u2014 ")
num("PRE-SUBMIT REFERENCE-DIFF. Before publishing, re-derive every number from its external source and diff "
    "it \u2014 run the reviewer's check yourself. Centralize shared identifiers (one revision constant) so a typo "
    "cannot diverge across files.", bold_lead="6 \u2014 ")
para("Why it survives human audit: every value points OUTWARD to a line a reviewer can open; the tests "
     "assert consequences and fail closed; observations carry a result record or are withheld. The reviewer "
     "checks FACTS and CITATIONS, not the agent's reading \u2014 which is exactly what a self-referenced artifact "
     "cannot offer.", GRAY)
para("Anti-patterns (each re-introduces self-reference): a green self-test as proof of correctness; "
     "re-guessing a shape/number after one is found wrong instead of capturing it from the reference; "
     "changing a dtype/label and declaring it fixed; a fail-OPEN gate that passes when the tracker/record is "
     "absent; citing a filename or model SHA as measurement provenance; and presenting an unsourced absolute "
     "as certified.", GRAY)

# ---- gates ----
h1("Gates")
for g in ["G0 — capacity / precision (every launch)",
          "G1 — coverage / lane (routes Thesis 1 vs 2 per op)",
          "Fusion roofline gate (bytes saved + AI lift + cache-resident; reference oracle where available)",
          "External-reference provenance gate (every op/shape/number/dtype/kernel/measurement cites an EXTERNAL authority or is EXPLICITLY-UNMODELED/withheld; tests assert consequences and fail closed; before any trust/perf-claim)",
          "Shape-provenance gate (every kernel's shapes CAPTURED from the real forward + human-signed manifest, before trust/perf-claims)",
          "Kernel-provenance gate (every authored kernel traced to the published reference: FAITHFUL/FRAGMENT(+enumerated gaps)/DIVERGENT + oracle + microbench-scope evidence recorded, REQUIRED before implementation review / re-authoring / perf-trust)",
          "Kernel\u2194cost-model reconciliation gate (every kernel maps to existing cost rows; every FRAGMENT gap is modeled-elsewhere / EXPLICITLY-UNMODELED / caller; dual paths named; enforced by a fail-closed self-test so the provenance and roofline cannot drift apart)",
          "Roofline-accounting gate (invariant self-test passes + one inventory + every quantity MODELED/MEASURED/EXPLICITLY-UNMODELED, before the report is published)",
          "Phase-A ceiling-or-surface (every op at its ceiling, else surface to user)",
          "G2 — validate-by-run (confirm the paper model in-engine)",
          "G2.5 — system-pathology (kernel-domination ≥ bar before per-op work)",
          "Accuracy gate · Human-gate (novel kernels)"]:
    bullet(g)

# ---- automated review-gated progression ----
h1("Automated review-gated progression — EXECUTOR \u2194 REVIEWER loop (optional orchestration)")
para("Each gate above can be driven as a two-agent loop: an EXECUTOR agent advances the gate and emits "
     "artifacts + a response; a DIFFERENT, independent REVIEWER agent audits them against the EXTERNAL "
     "reference and emits a machine-readable verdict; the orchestrator feeds a FAIL review back to the "
     "executor and iterates until PASS, a SURFACE-to-user signal, or a round cap. The user picks which model "
     "runs each role. Reference harness: tools/review_loop (agent-agnostic adapters; stdlib only).", bold=True)
bullet("SEPARATION OF ROLES: the reviewer must be a DIFFERENT agent/model that did NOT produce the "
       "artifacts \u2014 the whole point of external-reference provenance is an auditor who checks FACTS and "
       "CITATIONS, not the author's own reading. Same-agent self-review is self-consistency, not a gate.")
bullet("MACHINE-READABLE VERDICT CONTRACT: the reviewer emits one block \u2014 status PASS | FAIL | SURFACE, "
       "plus findings. FAIL -> findings fed back to the executor; SURFACE (or any unparseable/missing/"
       "malformed verdict) -> STOP and escalate to the user. FAIL-CLOSED: never treat 'no verdict' as pass.")
bullet("GUARDRAILS: a round cap (then SURFACE), commit between rounds for an audit trail, and a ledger of "
       "every round's verdict + findings. A gate that cannot converge is SURFACED with its history, not "
       "forced.")
bullet("SURFACE-TO-USER is a first-class outcome: when a finding needs a scope/trade-off decision the "
       "reference cannot settle, the executor or reviewer raises SURFACE and the loop halts for the user \u2014 "
       "exactly the points where a human call is required.")
para("The same disciplines apply inside the loop: the executor fixes the CLASS not the cited instance and "
     "changes CONSEQUENCES not labels; the reviewer runs independent counterexamples and passes only on "
     "reference-conformance within the declared partial scope. The pilot's 7-round roofline history "
     "(R\u2192F\u2192G\u2192H \u2192 PASS) replays through this harness as a regression of the loop itself.", GRAY)

# ---- platform card (example instantiation) ----
h1("Platform card — an INPUT, filled per target by the probe")
para("The method assumes NO specific values. Below is the pilot's instantiation (produced by the uArch "
     "probe); retarget by swapping the per-platform spec file. Every number here is measured, not assumed.", GRAY)
tbl = d.add_table(rows=1, cols=2)
tbl.style = "Light Grid Accent 1"
tbl.rows[0].cells[0].paragraphs[0].add_run("Property (pilot example)").bold = True
tbl.rows[0].cells[1].paragraphs[0].add_run("Value (DeepSeek-V4-Flash on Intel Xeon EMR, measured 2026-10-07)").bold = True
for k, v in [
    ("CPU", "Intel Xeon Platinum 8592+ (Emerald Rapids)"),
    ("Topology", "2 sockets × 64 cores; 1 NUMA domain per socket (SNC off)"),
    ("Per-domain RAM", "~503 GB (capacity floor)"),
    ("Memory BW (roofline ceiling)", "358.4 GB/s machine peak (8ch × DDR5-5600 × 8B); measured floor 214.4 GB/s = 60% of peak"),
    ("AMX bf16 (roofline ceiling)", "124.6 TFLOP/s machine peak (64c × 1024 FLOP/cyc × 1.9 GHz base); measured floor 40.3 TFLOP/s = 32% of peak"),
    ("Ridge (AI*)", "~348 FLOP/byte (machine peak)"),
    ("Optimization basis", "Optimize every op toward MACHINE PEAK; when levers are exhausted and it plateaus below peak, surface the residual gap and stop that op"),
    ("Implication", "Decode weight-streaming GEMMs are BW-bound across M∈{1,8,16,32,64}; only attention/sparse-attend are compute-bound"),
]:
    row = tbl.add_row().cells
    row[0].paragraphs[0].add_run(k)
    row[1].paragraphs[0].add_run(v)

# ---- method note ----
h1("Method note — all-ops-to-ceiling before wall-time ordering")
para("This playbook requires EVERY op at its roofline ceiling (Phase A) BEFORE wall-time / Amdahl "
     "ordering (Phase B), rather than ROI-skipping low-share ops up front. The inversion is deliberate: "
     "a cheap op can hide a systemic pathology, and completeness-first makes the Phase-B wall-time "
     "attribution trustworthy. Applies unchanged across models and platforms.", GRAY)

d.save(OUT)
print("SAVED:", OUT)
